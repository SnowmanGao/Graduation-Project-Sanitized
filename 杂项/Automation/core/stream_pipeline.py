import csv
import json
from copy import deepcopy
from pathlib import Path
from typing import Any, Callable, Optional

import openpyxl

from .paths import (
    DEFAULT_SPLIT_CSV_DIR,
    DEFAULT_FULL_CSV_PATH,
    DEFAULT_JSON_PATH,
    DEFAULT_MAP_PATH,
    DEFAULT_XLSX_PATH,
)


COMPONENT_ORDER = ["PTA", "EG", "PET", "DEG", "WATER", "BHET", "DOWTH-Q"]
COMPONENT_ROWS = set(COMPONENT_ORDER)
POLYMER_ROWS = {
    "PET",
    "DPN",
    "EFRAC",
    "T-EG",
    "B-EG",
    "T-TPA",
    "B-TPA",
    "T-DEG",
    "B-DEG",
    "MWN",
    "PDI",
    "ZMOM",
    "FMOM",
}
SUBSTRUCTURE_ROWS = {"T-EG", "B-EG", "T-TPA", "B-TPA", "T-DEG", "B-DEG"}
SCALAR_KEYS = {
    "相态",
    "温度",
    "压力",
    "摩尔汽相分率",
    "摩尔液相分率",
    "质量流量",
    "焓流量",
    "体积流量",
    "摩尔流量",
    "质量密度",
}
GROUP_KEYS = {"组分质量分率", "组分摩尔流量", "组分质量流量"}
SPLIT_TEMPLATE_FILES = [
    "流股数据0101.csv",
    "流股数据0102.csv",
    "流股数据02.csv",
    "流股数据03.csv",
]

type Pathlike = str | Path
type JsonData = dict[str, dict[str, Any]]

def as_path(value: Pathlike) -> Path:
    return value if isinstance(value, Path) else Path(value)


def ensure_parent(path: Pathlike) -> None:
    path = as_path(path)
    path.parent.mkdir(parents=True, exist_ok=True)


def normalize_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None
        if stripped.startswith("<"):
            return 0.0
        try:
            return float(stripped)
        except ValueError:
            return stripped
    return value


def parse_stream_workbook(
    xlsx_path: Pathlike, sheet_name: str | None = None
) -> JsonData:
    xlsx_path = as_path(xlsx_path)
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    assert worksheet is not None

    header_row = 4
    title_col = 3
    unit_col = 4
    stream_start_col = 5

    stream_names: list[str] = []
    col = stream_start_col
    while col <= worksheet.max_column:
        value = worksheet.cell(header_row, col).value
        if value is None:
            break
        stream_names.append(str(value).strip())
        col += 1

    if not stream_names:
        raise ValueError(f"[Error] 未在 {xlsx_path} 中找到流股名")

    result: JsonData = {
        stream_name: {} for stream_name in stream_names
    }
    current_mode: str | None = None

    for row in range(header_row + 1, worksheet.max_row + 1):
        title = worksheet.cell(row, title_col).value
        unit = worksheet.cell(row, unit_col).value

        if title is None:
            continue

        title = str(title).strip()

        if title == "MIXED子流股":
            current_mode = "base"
            continue
        if title == "质量分率":
            current_mode = "mass_frac"
            continue
        if title == "组分属性":
            current_mode = "polymer_attr"
            continue
        if title == "SFLOW":
            current_mode = "sflow"
            continue
        if title == "SFRAC":
            current_mode = "sfrac"
            continue

        for idx, stream_name in enumerate(stream_names, start=stream_start_col):
            value = normalize_value(worksheet.cell(row, idx).value)
            stream = result[stream_name]

            if title in SCALAR_KEYS:
                stream[title] = value
            elif 30 <= row <= 36 and title in COMPONENT_ROWS and unit == "kmol/hr":
                stream.setdefault("组分摩尔流量", {})[title] = value
            elif current_mode == "base" and title in COMPONENT_ROWS and unit == "kg/hr":
                stream.setdefault("组分质量流量", {})[title] = value
            elif (
                current_mode == "mass_frac" and title in COMPONENT_ROWS and unit is None
            ):
                stream.setdefault("组分质量分率", {})[title] = value
            elif current_mode == "polymer_attr" and title in POLYMER_ROWS:
                stream.setdefault("聚合物属性", {})[title] = value
            elif current_mode == "sflow" and title in SUBSTRUCTURE_ROWS:
                stream.setdefault("子结构流量", {})[title] = value
            elif current_mode == "sfrac" and title in SUBSTRUCTURE_ROWS:
                stream.setdefault("子结构分率", {})[title] = value

    return result


def transform_streams(data: JsonData) -> JsonData:
    for stream in data.values():
        stream.pop("子结构流量", None)
        stream.pop("子结构分率", None)
        polymer_attr = stream.pop("聚合物属性", {})
        stream["DPN (PET)"] = polymer_attr.get("DPN", 0) or 0
        stream.setdefault("组分质量流量", {}).pop("DOWTH-Q", None)
        stream.setdefault("组分质量分率", {}).pop("DOWTH-Q", None)
        stream.setdefault("组分摩尔流量", {}).pop("DOWTH-Q", None)

    return data


def write_json(data: dict[str, Any], path: Pathlike) -> None:
    path = as_path(path)
    ensure_parent(path)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_json(path: Pathlike) -> dict[str, Any]:
    path = as_path(path)
    return json.loads(path.read_text(encoding="utf-8"))


def hack_stream(
    data: JsonData,
    stream_name: str,
    hacker: Callable[[dict[str, Any]], None],
) -> JsonData:
    if stream_name not in data:
        raise KeyError(f"[Error] 找不到源流股: {stream_name}")
    
    try:
        hacker(data[stream_name])
    except Exception as err:
        print(f"[Error] 在 hack 流股 {stream_name} 时:\n{err}")

    return data


def hack_stream_based_on(
    data: JsonData,
    base_stream_name: str,
    target_stream_name: str,
    hacker: Optional[Callable[[dict[str, Any], dict[str, Any]], None]],
) -> JsonData:
    if base_stream_name not in data:
        raise KeyError(f"[Error] 找不到源流股: {base_stream_name}")
    
    data[target_stream_name] = deepcopy(data[base_stream_name])

    if hacker is None:
        return data
    
    try:
        hacker(data[base_stream_name], data[target_stream_name])
    except Exception as err:
        print(f"[Error] 在 hack 流股 {base_stream_name} -> {target_stream_name} 时:\n{err}")
    return data


def load_map_json(map_path: Pathlike) -> dict[str, str]:
    map_path = as_path(map_path)
    raw = load_json(map_path)
    if not isinstance(raw, dict):
        raise ValueError("[Error] map.json 顶层必须是对象")

    mapping: dict[str, str] = {}
    for stream_id, stream_name in raw.items():
        if not isinstance(stream_id, str) or not stream_id.strip():
            raise ValueError("[Error] map.json 中的编号键必须是非空字符串")
        if not isinstance(stream_name, str) or not stream_name.strip():
            raise ValueError(f"[Error] map.json 中编号 {stream_id!r} 的流股名必须是非空字符串")
        mapping[stream_id.strip()] = stream_name.strip()
    return mapping


def read_map_file(map_path: Pathlike) -> dict[str, list[str]]:
    id_to_name = load_map_json(map_path)
    mapping: dict[str, list[str]] = {}
    for stream_id, stream_name in id_to_name.items():
        mapping.setdefault(stream_name, []).append(stream_id)
    return mapping


def expand_streams_by_map(
    data: JsonData, name_to_ids: dict[str, list[str]]
) -> JsonData:
    expanded: JsonData = {}
    for stream_name, stream_ids in name_to_ids.items():
        if stream_name not in data:
            print(f"[Warn] 找不到 Aspen 流股 ID 映射目标: {stream_name}")
            continue
        for stream_id in stream_ids:
            if stream_id in expanded:
                raise ValueError(f"[Error] 检测到重复编号: {stream_id}")
            expanded[stream_id] = deepcopy(data[stream_name])
    return expanded


def build_fixed_id_dataset(
    data: JsonData, map_path: Pathlike
) -> JsonData:
    fixed: JsonData = {}
    for stream_name, stream_ids in read_map_file(map_path).items():
        payload = deepcopy(data.get(stream_name, {}))
        for stream_id in stream_ids:
            fixed[stream_id] = deepcopy(payload)
    return fixed


def get_sort_order() -> list[str]:
    return [
        "相态",
        "温度",
        "压力",
        "摩尔汽相分率",
        "摩尔液相分率",
        "质量密度",
        "焓流量",
        "体积流量",
        "质量流量",
        "DPN (PET)",
        "组分质量分率",
        "摩尔流量",
        "组分摩尔流量",
        "组分质量流量",
    ]


def get_unit_map() -> dict[str, str]:
    return {
        "相态": "",
        "温度": "C",
        "压力": "kPa",
        "摩尔汽相分率": "",
        "摩尔液相分率": "",
        "质量密度": "kg/cum",
        "焓流量": "kW",
        "体积流量": "cum/hr",
        "质量流量": "kg/hr",
        "DPN (PET)": "",
        "组分质量分率": "",
        "摩尔流量": "kmol/hr",
        "组分摩尔流量": "kmol/hr",
        "组分质量流量": "kg/hr",
    }


def build_csv_row_specs(
    data_by_id: JsonData, stream_ids: list[str] | None = None
) -> tuple[list[str], list[dict[str, str]]]:
    sort_order = get_sort_order()
    unit_map = get_unit_map()
    stream_ids = stream_ids or sorted(
        data_by_id.keys(), key=lambda value: (len(value), value)
    )
    row_specs: list[dict[str, str]] = []

    for key in sort_order:
        if not any(key in data_by_id.get(stream_id, {}) for stream_id in stream_ids):
            continue

        sample_value = next(
            data_by_id[stream_id][key]
            for stream_id in stream_ids
            if key in data_by_id.get(stream_id, {})
        )

        if isinstance(sample_value, dict):
            subkeys = set()
            for stream_id in stream_ids:
                value = data_by_id.get(stream_id, {}).get(key)
                if isinstance(value, dict):
                    subkeys.update(value.keys())

            ordered_subkeys = [item for item in COMPONENT_ORDER if item in subkeys]
            ordered_subkeys.extend(sorted(subkeys - set(ordered_subkeys)))

            row_specs.append(
                {
                    "type": "group_header",
                    "title": key,
                    "unit": unit_map.get(key, ""),
                    "group": key,
                }
            )
            for subkey in ordered_subkeys:
                row_specs.append(
                    {
                        "type": "group_item",
                        "title": subkey,
                        "unit": unit_map.get(key, ""),
                        "group": key,
                        "subkey": subkey,
                    }
                )
        else:
            row_specs.append(
                {
                    "type": "scalar",
                    "title": key,
                    "unit": unit_map.get(key, ""),
                    "group": key,
                }
            )

    return stream_ids, row_specs


def write_full_csv(data_by_id: JsonData, csv_path: Pathlike) -> None:
    csv_path = as_path(csv_path)
    stream_ids, row_specs = build_csv_row_specs(data_by_id)
    ensure_parent(csv_path)

    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["项目", "单位", *[f'="{stream_id}"' for stream_id in stream_ids]]
        )

        for spec in row_specs:
            row = [spec["title"], spec["unit"]]
            if spec["type"] == "group_header":
                row.extend([""] * len(stream_ids))
            elif spec["type"] == "scalar":
                row.extend(data_by_id.get(stream_id, {}).get(spec["group"]) for stream_id in stream_ids)  # type: ignore
            elif spec["type"] == "group_item":
                for stream_id in stream_ids:
                    value = data_by_id.get(stream_id, {}).get(spec["group"])
                    row.append(value.get(spec["subkey"]) if isinstance(value, dict) else None)  # type: ignore
            else:
                raise ValueError(f"[Error] 未知的 CSV 行类型: {spec['type']}")
            writer.writerow(row)


def format_cad_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".10g")
    return str(value)


def parse_stream_ids_from_header(header_row: list[str]) -> list[str]:
    stream_ids: list[str] = []
    for raw in header_row[2:]:
        cell = str(raw).strip()
        if not cell:
            stream_ids.append("")
            continue
        stream_ids.append(cell.lstrip("'").strip())
    return stream_ids


def fill_cad_template_rows(
    rows: list[list[str]], data_by_id: JsonData
) -> list[list[str]]:
    if not rows:
        return rows

    stream_ids = parse_stream_ids_from_header(rows[0])
    filled_rows: list[list[str]] = [rows[0]]
    current_group: str | None = None

    for row in rows[1:]:
        current = list(row)
        title = current[0].strip() if current else ""

        if title in GROUP_KEYS:
            current_group = title
            for idx in range(2, len(current)):
                current[idx] = ""
            filled_rows.append(current)
            continue

        if title not in COMPONENT_ROWS:
            current_group = None

        for col_idx, stream_id in enumerate(stream_ids, start=2):
            if col_idx >= len(current):
                current.extend([""] * (col_idx - len(current) + 1))

            stream = data_by_id.get(stream_id, {})
            if title in COMPONENT_ROWS and current_group:
                group_data = stream.get(current_group)
                value = group_data.get(title) if isinstance(group_data, dict) else None
            else:
                value = stream.get(title)
            current[col_idx] = format_cad_value(value)

        filled_rows.append(current)

    return filled_rows


def write_rows(path: Pathlike, rows: list[list[str]]) -> None:
    path = as_path(path)
    ensure_parent(path)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)


def read_rows(path: Pathlike) -> list[list[str]]:
    path = as_path(path)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def extract_streams(
    xlsx_path: Pathlike = DEFAULT_XLSX_PATH,
    output_path: Pathlike = DEFAULT_JSON_PATH,
    sheet_name: str | None = None,
) -> JsonData:
    data = transform_streams(parse_stream_workbook(xlsx_path, sheet_name=sheet_name))
    write_json(data, output_path)

    print("[Log] 已将", len(data), "流股数据解析并写入", Path(output_path).parts[-1])
    return data


def export_full_csv(
    input_path: Pathlike = DEFAULT_JSON_PATH,
    map_path: Pathlike = DEFAULT_MAP_PATH,
    output_path: Pathlike = DEFAULT_FULL_CSV_PATH,
) -> JsonData:
    data = load_json(input_path)
    expanded = expand_streams_by_map(data, read_map_file(map_path))
    write_full_csv(expanded, output_path)

    print("[Log] 已将精简后的 CSV 表格重新构建到", Path(output_path).parts[-1])
    print("[Log] 根据映射文件，最终得到", len(expanded), "条流股数据")
    return expanded


def export_split_csv_files(
    input_path: Pathlike = DEFAULT_JSON_PATH,
    map_path: Pathlike = DEFAULT_MAP_PATH,
    output_dir: Pathlike = DEFAULT_SPLIT_CSV_DIR,
) -> dict[str, Path]:
    data = load_json(input_path)
    data_by_id = build_fixed_id_dataset(data, map_path)
    output_dir = as_path(output_dir)

    written: dict[str, Path] = {}
    for filename in SPLIT_TEMPLATE_FILES:
        target_path = output_dir / filename
        template_rows = read_rows(target_path)
        filled_rows = fill_cad_template_rows(template_rows, data_by_id)
        write_rows(target_path, filled_rows)
        written[filename] = target_path

    print("[Log] 已导出", len(written), "个切分 CSV 文件（用于 CAD 动态链接）")
    return written


def run_all(
    xlsx_path: Pathlike = DEFAULT_XLSX_PATH,
    sheet_name: str | None = None,
    json_output_path: Pathlike = DEFAULT_JSON_PATH,
    map_path: Pathlike = DEFAULT_MAP_PATH,
    full_csv_output_path: Pathlike = DEFAULT_FULL_CSV_PATH,
    split_csv_output_dir: Pathlike = DEFAULT_SPLIT_CSV_DIR,
) -> dict[str, Any]:
    json_data = extract_streams(xlsx_path, json_output_path, sheet_name=sheet_name)
    expanded = export_full_csv(json_output_path, map_path, full_csv_output_path)
    split_files = export_split_csv_files(
        json_output_path, map_path, split_csv_output_dir
    )
    return {
        "json": json_data,
        "expanded": expanded,
        "cad_files": split_files,
    }
